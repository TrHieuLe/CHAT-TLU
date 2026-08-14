from __future__ import annotations

import csv
import logging
import os
import re
import shutil
import subprocess
from pathlib import Path
from typing import Protocol

import openpyxl

from app.core.config import settings

os.environ["CUDA_VISIBLE_DEVICES"] = settings.CUDA_VISIBLE_DEVICES

log = logging.getLogger(__name__)


class DocumentConverter(Protocol):
    def __call__(self, src: Path, dest: Path, *, to_console: bool = True) -> Path: ...


def _clean(text: str) -> str:
    text = re.sub(r"<br\s*/?>", " ", text)
    text = re.sub(r"<(.*?)>", r"[\1]", text)
    return text


def _find_libreoffice() -> str:
    env_path = os.getenv("LIBRE_OFFICE")
    if env_path and Path(env_path).exists():
        return env_path

    for cmd in ["libreoffice", "soffice"]:
        found = shutil.which(cmd)
        if found:
            return found

    raise RuntimeError(
        "Không tìm thấy LibreOffice/soffice. "
        "Cài bằng: sudo apt install libreoffice"
    )


def convert_to_pdf(src: Path, dest: Path, *, to_console: bool = True) -> Path:
    """
    Convert document (docx/doc, pptx) to pdf
    """
    if not src.exists():
        raise FileNotFoundError(src)

    office_cmd = _find_libreoffice()
    dest.parent.mkdir(parents=True, exist_ok=True)

    cmd = [
        office_cmd,
        "--headless",
        "--convert-to", "pdf",
        str(src),
        "--outdir", str(dest.parent),
    ]

    if to_console:
        log.info("Start converting %s ➜ %s", src.name, dest.name)

    result = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)

    if result.returncode != 0:
        raise RuntimeError(
            f"LibreOffice convert error:\nSTDOUT: {result.stdout}\nSTDERR: {result.stderr}"
        )

    converted = dest.parent / f"{src.stem}.pdf"
    if not converted.exists():
        raise FileNotFoundError(f"Không thấy file PDF sau khi convert: {converted}")

    if converted != dest:
        if dest.exists():
            dest.unlink()
        converted.rename(dest)

    log.info("✔ Pdf saved to %s", dest)
    return dest


def excel_to_csv(src: Path, dest_dir: Path, *, to_console: bool = True) -> list[Path]:
    """
    Convert Excel file to csv
    """
    if to_console:
        log.info(f"Start converting {src} to csv and saving to {dest_dir}")

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    paths = []

    for sheet in wb.sheetnames:
        ws = wb[sheet]
        safe_sheet_name = sheet.replace("/", "_").replace("\\", "_")
        csv_path = dest_dir / f"{safe_sheet_name} - {src.stem}.csv"

        with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                writer.writerow([cell if cell is not None else "" for cell in row])

        paths.append(csv_path)

        if to_console:
            log.info(f"✔ Sheet '{sheet}' saved to {csv_path}")

    return paths