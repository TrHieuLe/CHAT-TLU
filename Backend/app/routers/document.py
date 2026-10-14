import logging
import mimetypes
import unicodedata
import urllib.parse
import uuid
from pathlib import Path

import google.generativeai as genai
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, Form
from fastapi.responses import StreamingResponse, FileResponse
from pydantic import BaseModel
from qdrant_client import AsyncQdrantClient, QdrantClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.auth import require_admin
from app.core.config import settings
from app.models.database import ChatMessage, get_db, Document
from app.rag.qdrant_client_custom import qdrant_connection_kwargs
from app.rag.retriever import Retriever
from app.storage.helper import StorageHelper
from ingest import SUPPORTED, ingest_file

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api/document", tags=["document"])

_kw = qdrant_connection_kwargs()
_sync = QdrantClient(**_kw)
_async = AsyncQdrantClient(**_kw)
_retriever = Retriever(client=_sync, async_client=_async)

try:
    if settings.GEMINI_API_KEY or settings.GOOGLE_API_KEY:
        genai.configure(api_key=settings.EFFECTIVE_GEMINI_API_KEY)
except Exception as e:
    logger.warning("Chưa cấu hình API Key Gemini hợp lệ khi khởi động: %s", e)

APP_DIR = Path(__file__).resolve().parents[1]          # Backend/app
UPLOAD_DIR = APP_DIR / "data" / "images" / "uploads"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
DATA_DIR = Path(__file__).resolve().parents[2] / "data"  # Backend/data

storage_helper = StorageHelper()


def _nfc(text: str) -> str:
    return unicodedata.normalize("NFC", text or "").strip()


def _document_file_names(doc: Document) -> list[str]:
    """Các tên file vật lý có thể ứng với 1 Document (upload lưu original_name có đuôi, crawler thì không)."""
    names = [doc.original_name]
    if doc.filename and doc.file_type:
        names.append(f"{doc.filename}.{doc.file_type}")
    return [n for n in dict.fromkeys(names) if n]


async def resolve_document(
        db: AsyncSession, name: str, fuzzy: bool = True
) -> tuple[Path | None, Document | None]:
    """
    Tra cứu tài liệu theo tên mà frontend gửi lên (original_name, filename hoặc doc_name trong Qdrant):
    1. Tìm dòng Document khớp tên → thử các tên file vật lý của nó.
    2. Fallback: quét trực tiếp trên ổ đĩa.
    """
    name = _nfc(name)
    if not name:
        return None, None

    res = await db.execute(
        select(Document)
        .where(
            (Document.filename == name)
            | (Document.original_name == name)
            | (Document.filename == Path(name).stem)
        )
        .order_by(Document.created_at.desc())
        .limit(1)
    )
    document = res.scalars().first()

    if document:
        for candidate in _document_file_names(document):
            path = storage_helper.find_file(candidate, fuzzy=False)
            if path:
                return path, document

    return storage_helper.find_file(name, fuzzy=fuzzy), document


async def _resolve_request_name(db: AsyncSession, raw: str) -> tuple[Path | None, Document | None, str]:
    """Starlette đã decode path một lần; chỉ thử unquote thêm nếu tên gốc không tìm thấy."""
    name = _nfc(raw)
    path, document = await resolve_document(db, name)
    if not path:
        unquoted = _nfc(urllib.parse.unquote(raw))
        if unquoted != name:
            path, document = await resolve_document(db, unquoted)
            name = unquoted
    return path, document, name


@router.post("/upload", tags=["document"], dependencies=[Depends(require_admin)])
async def upload_document(
        db: AsyncSession = Depends(get_db),
        *,
        file: UploadFile = File(...),
        filename: str = Form(None),
        auto: bool = Form(True)
):
    logger.info("Start uploading document")
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in SUPPORTED:
        raise HTTPException(400, f"Định dạng {suffix or '(không có)'} không được hỗ trợ")
    if file.size is not None and file.size > settings.MAX_UPLOAD_MB * 1024 * 1024:
        raise HTTPException(413, f"File vượt quá {settings.MAX_UPLOAD_MB} MB")
    if not filename:
        filename = Path(file.filename).stem.strip()

    filename += Path(file.filename).suffix

    full_path = storage_helper.upload_v2(filename, file)
    if not full_path:
        raise HTTPException(500, "Đã có lỗi xảy ra khi lưu file, vui lòng thử lại sau")

    logger.info("Start ingest document")
    file_size = full_path.stat().st_size if full_path.exists() else 0
    file_type = full_path.suffix.lower().replace(".", "")

    doc = Document(
        id=str(uuid.uuid4()),
        filename=full_path.stem.strip(),
        original_name=full_path.name,
        file_type=file_type,
        file_size=file_size,
        status="processing",
    )

    try:
        doc.chunk_count = ingest_file(file_path=full_path, doc_id=doc.id, auto=auto)
        doc.status = "ok"
    except Exception as e:
        logger.error(f"Lỗi khi ingest file {filename}: {e}", exc_info=True)
        doc.status = "error"
        doc.error_msg = str(e)

    logger.info("Ingest document finished with status %s", doc.status)

    db.add(doc)
    await db.commit()
    await db.refresh(doc)

    return {
        "id": doc.id,
        "filename": doc.filename,
        "original_name": doc.original_name,
        "file_type": doc.file_type,
        "file_size": doc.file_size,
        "chunk_count": doc.chunk_count,
        "status": doc.status,
    }


@router.get("/list", tags=["document"], dependencies=[Depends(require_admin)])
async def list_documents(db: AsyncSession = Depends(get_db)):
    """Lấy danh sách tất cả tài liệu đã được tải lên và index."""
    res = await db.execute(
        select(Document).order_by(Document.created_at.desc())
    )
    documents = res.scalars().all()
    return [
        {
            "id": doc.id,
            "filename": doc.filename,
            "original_name": doc.original_name,
            "file_type": doc.file_type,
            "file_size": doc.file_size,
            "chunk_count": doc.chunk_count or 0,
            "status": doc.status,
            "created_at": doc.created_at.isoformat() if doc.created_at else None,
        }
        for doc in documents
    ]


@router.delete("/{doc_id}", tags=["document"], dependencies=[Depends(require_admin)])
async def delete_document(doc_id: str, db: AsyncSession = Depends(get_db)):
    """Xóa tài liệu khỏi cơ sở dữ liệu, vector store Qdrant và bộ lưu trữ file."""
    doc = await db.get(Document, doc_id)
    if not doc:
        raise HTTPException(status_code=404, detail="Không tìm thấy tài liệu")

    # Xóa khỏi Qdrant theo doc_name (doc_name = tên file có đuôi lúc ingest)
    from ingest import delete_doc_by_name
    doc_names = dict.fromkeys([*_document_file_names(doc), doc.filename])
    for doc_name in doc_names:
        if not doc_name:
            continue
        try:
            delete_doc_by_name(doc_name)
        except Exception as e:
            logger.warning("Không thể xóa vector Qdrant cho %s: %s", doc_name, e)

    # Xóa file vật lý khỏi storage (chỉ khớp chính xác)
    try:
        for candidate in _document_file_names(doc):
            if storage_helper.delete_v2(candidate):
                break
    except Exception as e:
        logger.warning("Không thể xóa file vật lý %s: %s", doc.original_name, e)

    await db.delete(doc)
    await db.commit()

    return {"ok": True, "message": f"Đã xóa tài liệu {doc.original_name} thành công"}


@router.get("/stats", tags=["document"], dependencies=[Depends(require_admin)])
async def get_document_stats(db: AsyncSession = Depends(get_db)):
    """Lấy thống kê tổng quan về tài liệu và chunks."""
    from sqlalchemy import func
    res = await db.execute(
        select(
            func.count(Document.id),
            func.coalesce(func.sum(Document.chunk_count), 0),
        )
    )
    total_docs, total_chunks = res.first() or (0, 0)

    return {
        "total_documents": total_docs,
        "total_chunks": total_chunks,
        "collection_name": settings.COLLECTION_NAME,
    }


@router.get("/reference/{filename:path}")
async def get_reference(db: AsyncSession = Depends(get_db), *, filename: str):
    """
    Phục vụ xem hoặc tải tài liệu tham khảo:
    1. Tìm trong DB Document table (theo filename hoặc original_name).
    2. Nếu không thấy hoặc file chưa đăng ký trong DB, quét đệ quy ổ đĩa trong Backend/data.
    3. Phục vụ với header Content-Disposition: inline và đúng MIME type (PDF, DOCX, TXT, MD...).
    """
    if not filename:
        raise HTTPException(400, "Tên nguồn tham khảo không được để trống")

    target_path, _document, decoded_filename = await _resolve_request_name(db, filename)
    # Tên tải về luôn giữ đúng đuôi file thật
    original_display_name = target_path.name if target_path else decoded_filename

    if not target_path or not target_path.is_file():
        logger.warning("Không tìm thấy file nguồn tham khảo: %s", decoded_filename)
        raise HTTPException(404, f"Nguồn tham khảo '{decoded_filename}' hiện không tồn tại trên hệ thống")

    try:
        content_type, _ = mimetypes.guess_type(target_path.name)
        if not content_type:
            ext = target_path.suffix.lower()
            if ext == ".docx":
                content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            elif ext == ".doc":
                content_type = "application/msword"
            elif ext in [".txt", ".md"]:
                content_type = "text/plain; charset=utf-8"
            elif ext == ".pdf":
                content_type = "application/pdf"
            else:
                content_type = "application/octet-stream"

        encoded_name = urllib.parse.quote(original_display_name)
        headers = {
            "Content-Disposition": f"inline; filename*=UTF-8''{encoded_name}",
        }

        return FileResponse(
            path=target_path,
            media_type=content_type,
            headers=headers,
        )

    except Exception as e:
        logger.error("Lỗi khi phục vụ file %s: %s", decoded_filename, e, exc_info=True)
        raise HTTPException(500, "Đã có lỗi xảy ra khi đọc tệp tin")


@router.get("/preview-text/{filename:path}")
async def get_preview_text(db: AsyncSession = Depends(get_db), *, filename: str):
    """
    Trích xuất và trả về nội dung text của tài liệu để xem nhanh trên giao diện modal.
    Hỗ trợ .txt, .md, .docx, .pdf, .xlsx/.csv.
    """
    if not filename:
        raise HTTPException(400, "Tên tài liệu không được để trống")

    target_path, document, decoded_filename = await _resolve_request_name(db, filename)

    if not target_path or not target_path.is_file():
        raise HTTPException(404, f"Không tìm thấy tài liệu '{decoded_filename}'")

    display_name = document.original_name if document and document.original_name else target_path.name
    ext = target_path.suffix.lower()
    file_size = target_path.stat().st_size
    text_content = ""

    try:
        if ext in [".txt", ".md", ".json"]:
            try:
                text_content = target_path.read_text(encoding="utf-8")
            except UnicodeDecodeError:
                text_content = target_path.read_text(encoding="cp1258", errors="ignore")

        elif ext in [".xlsx", ".xls", ".csv"]:
            try:
                from app.rag.table_extractor import extract_tables_from_excel
                excel_tables = extract_tables_from_excel(target_path)
                if excel_tables:
                    # table_text đã có sẵn tiêu đề "### Bảng: <sheet>"
                    text_content = "\n\n".join(
                        t.get("table_text", "")
                        for t in excel_tables
                        if t.get("table_text")
                    )
                elif ext == ".csv":
                    try:
                        text_content = target_path.read_text(encoding="utf-8")
                    except UnicodeDecodeError:
                        text_content = target_path.read_text(encoding="cp1258", errors="ignore")
                else:
                    text_content = f"Tệp bảng tính {ext.upper()} không có dữ liệu bảng hợp lệ."
            except Exception as e_xls:
                logger.warning("Excel preview extraction failed: %s", e_xls)
                text_content = f"Lỗi đọc bảng tính: {e_xls}"

        elif ext in [".docx", ".doc"]:
            try:
                import docx
                doc = docx.Document(target_path)
                paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
                text_content = "\n\n".join(paras)
            except Exception:
                try:
                    import docx2txt
                    text_content = docx2txt.process(str(target_path)) or ""
                except Exception:
                    text_content = ""

            # Bổ sung các bảng dữ liệu trích xuất dạng Markdown
            try:
                from app.rag.table_extractor import extract_tables_from_docx
                docx_tables = extract_tables_from_docx(target_path)
                if docx_tables:
                    table_mds = "\n\n".join(
                        f"### [Bảng dữ liệu trích xuất]\n{t.get('table_text', '')}"
                        for t in docx_tables
                        if t.get("table_text")
                    )
                    text_content = (text_content + "\n\n" + table_mds).strip()
            except Exception as e_tbl:
                logger.warning("Bóc tách bảng docx preview lỗi: %s", e_tbl)

        elif ext == ".pdf":
            try:
                import pdfplumber
                pages_text = []
                with pdfplumber.open(target_path) as pdf:
                    for idx, page in enumerate(pdf.pages[:15], 1):  # Giới hạn 15 trang đầu cho preview nhanh
                        p_text = page.extract_text() or ""
                        if p_text.strip():
                            pages_text.append(f"--- Trang {idx} ---\n{p_text}")
                text_content = "\n\n".join(pages_text)
            except Exception as e_pdf:
                logger.warning("pdfplumber preview extraction failed: %s", e_pdf)

        else:
            text_content = f"Tệp tin định dạng {ext.upper()} ({file_size} bytes). Vui lòng nhấn 'Tải về' để mở trên máy tính."

        return {
            "ok": True,
            "filename": target_path.name,
            "display_name": display_name,
            "file_type": ext.replace(".", ""),
            "file_size": file_size,
            "content": text_content.strip(),
        }

    except Exception as e:
        logger.error("Lỗi khi trích xuất text preview cho %s: %s", target_path.name, e)
        return {
            "ok": False,
            "filename": target_path.name,
            "display_name": display_name,
            "file_type": ext.replace(".", ""),
            "file_size": file_size,
            "content": f"Không thể trích xuất nội dung văn bản tự động: {e}",
        }


async def sync_local_data_documents(db: AsyncSession) -> int:
    """
    Quét đệ quy thư mục Backend/data/ và tự động thêm các tài liệu chưa có vào bảng Document.
    """
    data_dir = DATA_DIR
    if not data_dir.exists():
        return 0

    added_count = 0
    supported_exts = {".pdf", ".docx", ".doc", ".txt", ".md", ".xlsx", ".xls", ".csv"}

    for file_path in data_dir.rglob("*"):
        if not file_path.is_file():
            continue
        if file_path.suffix.lower() not in supported_exts:
            continue
        # Bỏ qua ảnh upload và bài cào web (crawler tự quản lý dòng Document)
        rel_parts = file_path.relative_to(data_dir).parts
        if {"uploads", "images", "crawled"} & set(rel_parts):
            continue

        filename_stem = file_path.stem.strip()
        original_name = file_path.name

        # Kiểm tra xem đã có trong DB chưa
        res = await db.execute(
            select(Document).where(
                (Document.original_name == original_name) | (Document.filename == filename_stem)
            )
        )
        existing = res.scalars().first()

        if not existing:
            doc = Document(
                id=str(uuid.uuid4()),
                filename=filename_stem,
                original_name=original_name,
                file_type=file_path.suffix.lower().replace(".", ""),
                file_size=file_path.stat().st_size,
                chunk_count=0,
                status="ok",
            )
            db.add(doc)
            added_count += 1

    if added_count > 0:
        await db.commit()
        logger.info("✅ Đã tự động đồng bộ %d tài liệu từ data/ vào cơ sở dữ liệu", added_count)

    return added_count
